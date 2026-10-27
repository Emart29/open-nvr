# Copyright (c) 2026 OpenNVR
# Licensed under the GNU Affero General Public License v3.0 (AGPL-3.0)
"""Camera import and export (services/camera_io.py, services/camera_import.py).

Pinned here:

* CSV and Excel read the same; headers match loosely; blank rows drop out;
* an NVR row expands to one camera per channel, with the vendor's RTSP
  pattern or an rtsp_url holding {ch};
* every row is checked on its own, and a bad row never blocks the rest;
* matching: a row with an RTSP URL matches by URL (logins ignored), one with
  only an address matches the one camera there; importing the same file
  again changes nothing; a camera the importer may not change is refused;
* export leaves passwords, and the logins inside RTSP URLs, out unless asked,
  which needs MFA; a cell that would run as a formula is written as text;
* the whole round trip through the API: preview, import, poll, re-import.

Run with:
    cd server && pytest tests/test_camera_io.py -v
"""
from __future__ import annotations

import datetime as _dt

if not hasattr(_dt, "UTC"):
    _dt.UTC = _dt.timezone.utc  # noqa: UP017

import io
import os
import secrets
import sys
import time
import types as _types
from pathlib import Path

import pytest

REPO_ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(REPO_ROOT / "server"))

from cryptography.fernet import Fernet  # noqa: E402

os.environ.setdefault("DATABASE_URL", "postgresql://u:p@localhost/x")
os.environ.setdefault("SECRET_KEY", secrets.token_urlsafe(48))
os.environ.setdefault("MEDIAMTX_SECRET", secrets.token_hex(32))
os.environ.setdefault("INTERNAL_API_KEY", secrets.token_urlsafe(48))
os.environ.setdefault("CREDENTIAL_ENCRYPTION_KEY", Fernet.generate_key().decode())

_lm = _types.ModuleType("core.logging_config")


class _L:
    def __getattr__(self, _n):
        return lambda *a, **k: None


_lm.__getattr__ = lambda _n: _L()
_lm.setup_logging = lambda *a, **k: None
sys.modules.setdefault("core.logging_config", _lm)

from fastapi import FastAPI  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402
from sqlalchemy import create_engine  # noqa: E402
from sqlalchemy.orm import sessionmaker  # noqa: E402

import core.auth as core_auth  # noqa: E402
import core.database as core_database  # noqa: E402
from core.database import Base, get_db  # noqa: E402
from models import AuditLog, Camera, CameraPermission, Permission, Role, RolePermission, User  # noqa: E402
from routers import cameras as cameras_router  # noqa: E402
from services import camera_io as cio  # noqa: E402

CSV_HEADER = "name,location,ip_address,port,username,password,rtsp_url,substream_url,description,vlan,active,channels,vendor\n"


def _rows(text: str):
    return cio.read_table(text.encode(), "cams.csv")


def _plan(text: str):
    return cio.plan_rows(_rows(text))


# ── Reading ──────────────────────────────────────────────────────────────


def test_headers_match_loosely_and_blank_rows_drop():
    rows = _rows("Camera Name;IP Address;Zone;Whatever\nGate;10.0.0.5;North;x\n;;;\nDoor;10.0.0.6;;\n")
    assert [(n, r["name"], r["ip_address"], r.get("location")) for n, r in rows] == [
        (2, "Gate", "10.0.0.5", "North"), (4, "Door", "10.0.0.6", "")]


def test_a_file_without_a_name_column_is_refused():
    with pytest.raises(cio.ImportFileError, match="name"):
        _rows("ip,port\n10.0.0.1,554\n")


def test_excel_reads_like_csv():
    body = cio.to_xlsx([["Gate", "North", "10.0.0.5", 554, "admin", "", "", "", "", "", "yes", "", ""]])
    rows = cio.read_table(body, "cams.xlsx")
    assert rows[0][1]["name"] == "Gate" and rows[0][1]["port"] == "554" and rows[0][1]["active"] == "yes"


def test_a_quoted_formula_cell_reads_back_as_written():
    body = cio.to_csv([["=HYPERLINK(1)", "", "10.0.0.5", "", "", "", "", "", "", "", "", "", ""]])
    assert b"'=HYPERLINK" in body
    assert cio.read_table(body, "x.csv")[0][1]["name"] == "=HYPERLINK(1)"


def test_excel_writes_text_cells_never_formulas():
    from openpyxl import load_workbook

    body = cio.to_xlsx([["=1+1", "", "10.0.0.5", "", "", "", "", "", "", "", "", "", ""]])
    ws = load_workbook(io.BytesIO(body)).worksheets[0]
    assert ws["A2"].value == "=1+1" and ws["A2"].data_type == "s"


# ── Checking and expanding ────────────────────────────────────────────────


def test_an_nvr_row_becomes_a_camera_per_channel():
    planned = _plan(CSV_HEADER + "NVR North,North,10.1.0.20,,admin,pw,,,,,yes,1-3,hikvision\n")
    assert [p.fields["name"] for p in planned] == ["NVR North Ch 1", "NVR North Ch 2", "NVR North Ch 3"]
    assert planned[1].fields["rtsp_url"] == "rtsp://10.1.0.20:554/Streaming/Channels/201"
    assert planned[1].fields["substream_url"] == "rtsp://10.1.0.20:554/Streaming/Channels/202"
    assert all(p.fields["location"] == "North" and p.status == "new" for p in planned)


def test_dahua_channels_and_a_custom_pattern():
    dahua = _plan(CSV_HEADER + "DVR,,10.1.0.21,37777,,,,,,,,2,cp plus\n")
    assert dahua[0].fields["rtsp_url"] == "rtsp://10.1.0.21:37777/cam/realmonitor?channel=2&subtype=0"
    custom = _plan(CSV_HEADER + "Box,,,,,,rtsp://10.1.0.22:8554/ch{ch}/main,,,,,\"1,4\",\n")
    assert [p.fields["rtsp_url"] for p in custom] == ["rtsp://10.1.0.22:8554/ch1/main", "rtsp://10.1.0.22:8554/ch4/main"]
    assert custom[0].fields["ip_address"] == "10.1.0.22" and custom[0].fields["port"] == 8554


@pytest.mark.parametrize("line, wanted", [
    (",,10.0.0.1,,,,,,,,,,", "name is missing"),
    ("Cam,,,,,,,,,,,,", "ip_address is missing"),
    ("Cam,,not an ip!,,,,,,,,,,", "is not an address"),
    ("Cam,,10.0.0.1,99999,,,,,,,,,", "port"),
    ("Cam,,10.0.0.1,,,,http://x/y,,,,,,", "rtsp://"),
    ("Cam,,10.0.0.1,,,,,,,,maybe,,", "yes or no"),
    ("Cam,,10.0.0.1,,,,,,,,,1-4,", "needs a vendor"),
    ("Cam,,10.0.0.1,,,,,,,,,1-4,sony", "is not one of"),
    ("Cam,,10.0.0.1,,,,,,,,,9-2,hikvision", "channels"),
    ("Cam,,10.0.0.1,,,,rtsp://10.0.0.1/s{ch},,,,,,", "no channels"),
])
def test_a_bad_row_says_why(line, wanted):
    (p,) = _plan(CSV_HEADER + line + "\n")
    assert p.status == "error" and any(wanted in e for e in p.errors), p.errors


def test_a_bad_row_does_not_stop_the_others():
    planned = _plan(CSV_HEADER + ",North,,,,,,,,,,,\nGood,,10.0.0.9,,,,,,,,,,\n")
    assert [p.status for p in planned] == ["error", "new"]


def test_channels_parse():
    assert cio.parse_channels("1-3, 7 ,5") == [1, 2, 3, 5, 7]
    with pytest.raises(ValueError):
        cio.parse_channels("1-300")


# ── Matching ───────────────────────────────────────────────────────────────


def _cam(id, name, ip, rtsp=None, **kw):
    return _types.SimpleNamespace(id=id, name=name, ip_address=ip, rtsp_url=rtsp, port=kw.pop("port", 554),
                                  location=None, description=None, vlan=None, username=None, password=None,
                                  substream_url=None, is_active=True, **kw)


def test_matching_by_url_ignores_logins_and_reimport_is_unchanged():
    existing = [_cam(1, "NVR Ch 1", "10.1.0.20", "rtsp://admin:x@10.1.0.20/Streaming/Channels/101")]
    planned = _plan(CSV_HEADER + "NVR,,10.1.0.20,,,,,,,,,1-2,hikvision\n")
    cio.match_existing(planned, existing, None)
    # Ch 1 is there under the same name; its substream would be added.
    assert [(p.status, p.camera_id) for p in planned] == [("update", 1), ("new", None)]
    assert planned[0].changes == ["substream_url"]


def test_matching_by_address_needs_it_to_be_unambiguous():
    existing = [_cam(1, "A", "10.0.0.5"), _cam(2, "B", "10.0.0.6"), _cam(3, "C", "10.0.0.6")]
    planned = _plan(CSV_HEADER + "A,North,10.0.0.5,,,,,,,,,,\nB,,10.0.0.6,,,,,,,,,,\nNew,,10.0.0.7,,,,,,,,,,\n")
    cio.match_existing(planned, existing, None)
    assert [p.status for p in planned] == ["update", "error", "new"]
    assert planned[0].changes == ["location"]
    assert "give the rtsp_url" in planned[1].errors[0]


def test_twice_in_the_file_is_a_duplicate():
    planned = _plan(CSV_HEADER + "A,,,,,,rtsp://10.0.0.5/s,,,,,,\nA again,,,,,,rtsp://x:y@10.0.0.5:554/s,,,,,,\n")
    cio.match_existing(planned, [], None)
    assert [p.status for p in planned] == ["new", "duplicate"]


def test_a_camera_the_importer_may_not_change_is_refused():
    existing = [_cam(1, "A", "10.0.0.5")]
    planned = _plan(CSV_HEADER + "A,North,10.0.0.5,,,,,,,,,,\n")
    cio.match_existing(planned, existing, manageable=set())
    assert planned[0].status == "error" and "may not change" in planned[0].errors[0]


def test_a_blank_cell_never_changes_a_camera():
    existing = [_cam(1, "A", "10.0.0.5", port=8554)]
    existing[0].location = "North"
    planned = _plan(CSV_HEADER + "A,,10.0.0.5,,,,,,,,,,\n")
    cio.match_existing(planned, existing, None)
    assert planned[0].status == "unchanged"


# ── Export ───────────────────────────────────────────────────────────────


def test_export_leaves_logins_out_unless_asked():
    cam = _cam(1, "A", "10.0.0.5", "rtsp://admin:s3cret@10.0.0.5:554/cam/realmonitor?channel=1&subtype=0")
    cam.username, cam.password = "admin", "s3cret"
    plain = cio.export_rows([cam], include_passwords=False)[0]
    assert "s3cret" not in str(plain)
    assert plain[6] == "rtsp://10.0.0.5:554/cam/realmonitor?channel=1&subtype=0"
    full = cio.export_rows([cam], include_passwords=True)[0]
    assert full[5] == "s3cret" and "admin:s3cret@" in full[6]


def test_an_export_reads_back_as_an_unchanged_import():
    cam = _cam(1, "Gate", "10.0.0.5", "rtsp://10.0.0.5:554/s")
    cam.location = "North"
    body = cio.to_csv(cio.export_rows([cam], include_passwords=False))
    planned = cio.plan_rows(cio.read_table(body, "export.csv"))
    cio.match_existing(planned, [cam], None)
    assert planned[0].status == "unchanged"


# ── Through the API ─────────────────────────────────────────────────────────


@pytest.fixture
def env(monkeypatch, tmp_path):
    monkeypatch.setattr(core_auth, "auth_logger", _L(), raising=False)
    monkeypatch.setattr(cameras_router, "camera_logger", _L(), raising=False)
    # A file, not sqlite:// on one shared connection: the import job writes
    # from the event loop while the test polls from its own thread, and on a
    # single connection one session's rollback lands in the other's commit.
    eng = create_engine(f"sqlite:///{tmp_path / 'cams.db'}", connect_args={"check_same_thread": False})
    Base.metadata.create_all(eng)
    Session = sessionmaker(bind=eng, expire_on_commit=False)
    # The import job opens its own sessions: point them at this database.
    monkeypatch.setattr(core_database, "SessionLocal", Session)

    db = Session()
    manage = Permission(name="cameras.manage", description="")
    role = Role(name="operator", description="")
    db.add_all([manage, role])
    db.flush()
    db.add(RolePermission(role_id=role.id, permission_id=manage.id))
    admin = User(username="admin", email="a@x", hashed_password="x", is_active=True,
                 is_superuser=True, password_set=True, role_id=role.id)
    operator = User(username="op", email="o@x", hashed_password="x", is_active=True,
                    is_superuser=False, password_set=True, role_id=role.id)
    db.add_all([admin, operator])
    db.commit()
    for u in (admin, operator):
        db.refresh(u)
        _ = [p.name for p in (u.role.permissions or [])]
        db.expunge(u)
    db.close()

    # Creating a camera provisions its stream: keep MediaMTX out of it.
    import services.camera_service as cs

    async def _create(db, camera_create, owner_id):
        cam = Camera(name=camera_create.name, ip_address=camera_create.ip_address, port=camera_create.port,
                     rtsp_url=camera_create.rtsp_url, substream_url=camera_create.substream_url,
                     location=camera_create.location, username=camera_create.username,
                     owner_id=owner_id, is_active=True)
        cam.password = camera_create.password
        db.add(cam)
        db.commit()
        db.refresh(cam)
        return cam

    monkeypatch.setattr(cs.CameraService, "create_camera", staticmethod(_create))

    app = FastAPI()
    app.include_router(cameras_router.router, prefix="/api/v1")

    def _db():
        s = Session()
        try:
            yield s
        finally:
            s.close()

    current = {"user": admin}
    app.dependency_overrides[get_db] = _db
    app.dependency_overrides[core_auth.get_current_active_user] = lambda: current["user"]
    with TestClient(app) as tc:
        yield _types.SimpleNamespace(tc=tc, current=current, admin=admin, operator=operator, Session=Session)


def _upload(text: str, name="cams.csv"):
    return {"file": (name, text.encode(), "text/csv")}


def _run(env, text: str, **form):
    r = env.tc.post("/api/v1/cameras/import", files=_upload(text), data=form)
    assert r.status_code == 200, r.text
    job = r.json()
    for _ in range(200):
        job = env.tc.get(f"/api/v1/cameras/import/jobs/{job['id']}").json()
        if job["state"] == "done":
            return job
        time.sleep(0.05)
    raise AssertionError("the import did not finish")


FILE = CSV_HEADER + (
    "Gate,North,192.168.1.64,,admin,pw1,rtsp://192.168.1.64:554/live,,,,yes,,\n"
    "NVR South,South,192.168.2.20,,admin,pw2,,,,,yes,1-3,hikvision\n"
    ",,,,,,,,,,,,\n"
    "Broken,,,,,,,,,,,,\n"
)


def test_preview_saves_nothing_and_says_what_would_happen(env):
    r = env.tc.post("/api/v1/cameras/import/preview", files=_upload(FILE))
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["summary"] == {"total": 5, "new": 4, "update": 0, "unchanged": 0, "duplicate": 0, "error": 1}
    assert not any("pw" in str(row) for row in body["rows"])
    s = env.Session()
    assert s.query(Camera).count() == 0


def test_import_then_reimport_changes_nothing(env):
    job = _run(env, FILE)
    assert job["counts"] == {"created": 4, "skipped": 1}
    s = env.Session()
    names = sorted(c.name for c in s.query(Camera).all())
    assert names == ["Gate", "NVR South Ch 1", "NVR South Ch 2", "NVR South Ch 3"]
    assert s.query(Camera).filter_by(name="NVR South Ch 2").one().rtsp_url.endswith("/Streaming/Channels/201")
    assert s.query(AuditLog).filter_by(action="camera.import").count() == 1
    s.close()

    again = _run(env, FILE, mode="add_update")
    assert again["counts"] == {"unchanged": 4, "skipped": 1}


def test_updates_only_when_asked(env):
    _run(env, FILE)
    moved = FILE.replace("Gate,North", "Gate,East")
    assert _run(env, moved)["counts"].get("skipped") == 2
    assert _run(env, moved, mode="add_update")["counts"].get("updated") == 1
    s = env.Session()
    assert s.query(Camera).filter_by(name="Gate").one().location == "East"


def test_an_operator_cannot_change_someone_elses_camera(env):
    _run(env, FILE)
    env.current["user"] = env.operator
    body = env.tc.post("/api/v1/cameras/import/preview", files=_upload(FILE.replace("Gate,North", "Gate,East"))).json()
    gate = next(r for r in body["rows"] if r["name"] == "Gate")
    assert gate["status"] == "error" and "may not change" in gate["errors"][0]


def test_a_bad_file_is_a_422(env):
    r = env.tc.post("/api/v1/cameras/import/preview", files=_upload("ip,port\n1.2.3.4,554\n"))
    assert r.status_code == 422 and "name" in r.json()["detail"]


def test_template_downloads_in_both_formats(env):
    r = env.tc.get("/api/v1/cameras/import/template", params={"format": "csv"})
    assert r.status_code == 200 and r.content.startswith("﻿name,".encode())
    r = env.tc.get("/api/v1/cameras/import/template")
    assert r.status_code == 200 and r.headers["content-type"].startswith("application/vnd.openxml")
    planned = cio.plan_rows(cio.read_table(r.content, "t.xlsx"))
    assert len(planned) == 18  # 2 cameras and a 16-channel NVR


def test_export_needs_mfa_for_passwords_and_is_audited(env):
    _run(env, FILE)
    r = env.tc.get("/api/v1/cameras/export", params={"format": "csv"})
    assert r.status_code == 200 and b"pw1" not in r.content and b"Gate" in r.content
    r = env.tc.get("/api/v1/cameras/export", params={"include_passwords": True})
    assert r.status_code == 400  # no MFA set up on this account
    r = env.tc.get("/api/v1/cameras/export", params={"location": "south", "format": "xlsx"})
    planned = cio.plan_rows(cio.read_table(r.content, "e.xlsx"))
    assert [p.fields["name"] for p in planned] == ["NVR South Ch 1", "NVR South Ch 2", "NVR South Ch 3"]
    s = env.Session()
    assert s.query(AuditLog).filter_by(action="camera.export").count() == 2


def test_export_covers_only_cameras_the_caller_manages(env):
    _run(env, FILE)
    s = env.Session()
    gate = s.query(Camera).filter_by(name="Gate").one()
    s.add(CameraPermission(user_id=env.operator.id, camera_id=gate.id, can_view=True, can_manage=True))
    s.commit()
    env.current["user"] = env.operator
    body = env.tc.get("/api/v1/cameras/export").content.decode("utf-8-sig")
    assert "Gate" in body and "NVR South" not in body
